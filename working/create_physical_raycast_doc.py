from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Huong_dan_Ray_Casting_vat_ly.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    node = OxmlElement("w:keepNext")
    p_pr.append(node)


def set_font(run, name="Aptos", size=10.5, bold=False, color="222222", italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(34, 34, 34)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in (("Title", 25, "000000"), ("Heading 1", 17, "000000"), ("Heading 2", 13, "000000"), ("Heading 3", 11, "000000")):
        st = styles[name]
        st.font.name = "Aptos Display" if name in ("Title", "Heading 1") else "Aptos"
        st._element.rPr.rFonts.set(qn("w:ascii"), st.font.name)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), st.font.name)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        st.paragraph_format.space_after = Pt(6)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "LAB SAM  |  Ray Casting vật lý"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        set_font(r, size=8.5, color="666666")
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Hướng dẫn thực hành  |  ")
    set_font(r, size=8.5, color="666666")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)


def add_p(doc, text="", bold_prefix=None, style=None, italic=False):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2, italic=italic)
    else:
        r = p.add_run(text)
        set_font(r, italic=italic)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        r = p.add_run(item)
        set_font(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_font(r)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.right_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F5F7")
    pPr.append(shd)
    r = p.add_run(text)
    set_font(r, name="Consolas", size=9, color="222222")
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, size=9, bold=True, color="FFFFFF")
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx % 2 == 1:
                set_cell_shading(cell, "F5F8FB")
            p = cell.paragraphs[0]
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_font(r, size=9)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_checklist(doc, items):
    rows = [["☐", item] for item in items]
    add_table(doc, ["Xong", "Hạng mục kiểm tra"], rows, widths=[1.2, 15.8])


def main():
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run("Hướng dẫn Ray Casting vật lý")
    set_font(r, name="Aptos Display", size=25, bold=True, color="000000")
    sub = doc.add_paragraph()
    r = sub.add_run("Từ camera iPhone đến tọa độ đám cháy trong hệ mét của phòng")
    set_font(r, size=13, color="4F6475")
    add_p(doc, "Tài liệu này mô tả quy trình thực hành cho nhánh Ray Casting vật lý của dự án LAB SAM. Mục tiêu là tạo một thí nghiệm phòng nhỏ có camera cố định, calibration đo thật, mesh cùng hệ tọa độ và ground truth XYZ cho LED hoặc marker an toàn. Kết quả cuối phải phân biệt rõ giữa kiểm thử hình học và độ chính xác vật lý trong phòng.")
    add_p(doc, "Phạm vi: phòng khách hoặc một vùng phòng được giới hạn và đo đạc; không cần dùng lửa thật. LED/marker được xem là mục tiêu thay thế để xác định điểm cần chiếu tia.")

    add_heading(doc, "Thông số đã chốt", 1)
    add_table(doc, ["Thành phần", "Cấu hình bắt buộc", "Ghi chú"], [
        ["Camera", "iPhone 13 Pro Max, lens chính 1x, cố định trên tripod", "Khóa focus, exposure và zoom; không đổi vị trí sau calibration"],
        ["Vùng test", "Khoảng 3–4 m rộng và 4–6 m sâu", "Bắt đầu bằng vùng trống trong phòng khách"],
        ["ArUco", "DICT_6X6_1000, ID 0–7", "Không trộn dictionary khác trong cùng session"],
        ["ArUco kích thước", "Cạnh ô đen 120 mm, lề trắng 20 mm", "In PDF/SVG ở 100% actual size; đo lại sau in"],
        ["Checkerboard", "9 × 6 inner corners, 10 × 7 squares", "Ô 25 mm; bảng 250 × 175 mm; OpenCV dùng (9, 6)"],
        ["Hệ tọa độ", "X sang phải, Y đi vào phòng, Z hướng lên", "Gốc đặt ở góc sau trái vùng test; đơn vị mét"],
        ["Marker phòng", "Tối thiểu 6 nhìn thấy, khuyến nghị 6–10", "Phân bố rộng, có độ cao khác nhau nếu có thể"],
        ["Mục tiêu", "LED hoặc marker màu; không cần lửa thật", "Quy ước điểm là tâm LED hoặc điểm đáy, chọn một và giữ nguyên"],
        ["Ray Casting", "Single, 3, 5 và 7 rays; ưu tiên cross/square", "Ray cắt mesh theo calibration thật"],
        ["Lọc", "Outlier rejection rồi EMA hoặc EKF", "Không dùng lọc để che lỗi calibration"],
        ["Đầu ra", "XYZ, MAE/median/P95, hit rate, latency, PNG/HTML/PLY", "Chỉ gọi là physical accuracy khi có dữ liệu đo thật"],
    ], widths=[3.2, 6.0, 7.8])

    add_heading(doc, "Luồng tổng thể", 1)
    add_code(doc, "iPhone cố định\n    -> checkerboard -> K và distortion\n    -> ArUco 2D + tọa độ góc 3D -> solvePnP -> camera pose\n    -> đo phòng -> room mesh theo mét\n    -> LED tại XYZ đã đo -> p_fire_pixel và fire_xyz_world\n    -> detector -> ROI mixed -> undistortion\n    -> multi-ray mesh intersection -> outlier rejection\n    -> EMA hoặc EKF -> benchmark -> PNG / HTML / PLY")
    add_p(doc, "Ba thành phần camera calibration, room mesh và ground truth XYZ phải dùng cùng một hệ tọa độ và cùng đơn vị mét. Nếu thiếu một thành phần, pipeline vẫn có thể chạy nhưng kết quả chỉ là provisional hoặc synthetic.")

    doc.add_page_break()
    add_heading(doc, "Chuẩn bị thiết bị và vật tư", 1)
    add_p(doc, "Đã có tripod và đèn LED. Các mục còn lại cần chuẩn bị trước khi chụp dữ liệu.")
    add_table(doc, ["Hạng mục", "Số lượng khuyến nghị", "Mục đích"], [
        ["iPhone 13 Pro Max", "1", "Nguồn ảnh/video; giữ cố định trên tripod"],
        ["Checkerboard in", "1 bảng", "Calibrate K và distortion"],
        ["ArUco in", "8 marker ID 0–7", "Xác định pose camera trong hệ phòng"],
        ["Bìa cứng hoặc foam board", "9 miếng", "Làm phẳng checkerboard và marker"],
        ["Thước dây", "1", "Đo phòng, marker, LED và vật cản"],
        ["Thước thẳng", "1", "Kiểm tra kích thước in và cạnh marker"],
        ["Băng keo giấy", "1 cuộn", "Đánh dấu gốc, trục và vị trí mục tiêu"],
        ["Máy tính", "1", "Lưu ảnh và chạy .venv của dự án"],
    ], widths=[5.0, 4.0, 8.0])
    add_heading(doc, "In checkerboard", 2)
    add_p(doc, "Dùng file checkerboard_9x6_inner_25mm.svg trong thư mục calibration_targets/aruco_svg_printables. Khi in chọn A4, Actual size hoặc 100%, không chọn Fit to page. Sau khi in, đo lại ô vuông; nếu kích thước thực tế khác 25 mm, phải ghi lại giá trị thật.")
    add_heading(doc, "In marker ArUco", 2)
    add_p(doc, "Dùng tám file aruco_dict_6x6_1000_id_00.svg đến aruco_dict_6x6_1000_id_07.svg. In mỗi marker trên một tờ A4, giữ lề trắng, dán lên bìa cứng và đo cạnh đen. Không cắt sát vào ô đen và không dùng ảnh chụp màn hình để in.")
    add_checklist(doc, [
        "Đã in checkerboard với 9 × 6 inner corners và ô 25 mm",
        "Đã in đủ ArUco ID 0–7 bằng DICT_6X6_1000",
        "Đã đo lại cạnh đen 120 mm của marker",
        "Checkerboard và marker đã dán lên mặt phẳng cứng",
        "Đã chuẩn bị thước dây, thước thẳng và băng keo",
        "Đã chọn vùng test và loại bỏ vật thể di động",
    ])

    add_heading(doc, "Chọn và bố trí phòng", 1)
    add_p(doc, "Phòng khách có thể dùng, nhưng chỉ nên chọn một vùng test hẹp. Không dùng toàn bộ phòng khách nếu chưa đo được hình học. Giai đoạn đầu nên dọn vùng sàn, sau đó mới thêm bàn hoặc tủ để đánh giá occlusion.")
    add_table(doc, ["Giai đoạn", "Cấu hình", "Mục tiêu"], [
        ["1 Phòng trống", "Mesh chỉ có sàn, LED đặt trên sàn", "Kiểm tra calibration và ray-surface intersection"],
        ["2 Có vật cản", "Thêm một bàn hoặc tủ đã đo", "Kiểm tra mặt phẳng và occlusion"],
        ["3 Gần thực tế", "Giữ một số đồ đạc cố định", "Đánh giá độ ổn định trong phòng khách"],
    ], widths=[4.0, 7.0, 6.0])
    add_p(doc, "Kích thước khởi đầu khuyến nghị là 3–4 m theo chiều ngang và 4–6 m theo chiều sâu. Camera nên cao khoảng 1.5–2.0 m, nhìn được LED ở vùng gần, giữa và xa, đồng thời vẫn nhìn thấy tối thiểu sáu marker.")

    doc.add_page_break()
    add_heading(doc, "Định nghĩa hệ tọa độ phòng", 1)
    add_p(doc, "Chọn hệ tọa độ trước khi dán marker. Một quy ước dễ dùng là gốc ở góc sau trái của vùng test, X sang phải, Y đi vào phòng và Z hướng lên. Đánh dấu gốc O, trục X và trục Y bằng băng keo trên sàn.")
    add_code(doc, "O = góc sau trái vùng test\nX = sang phải\nY = đi vào phòng\nZ = hướng lên\nunits = metres")
    add_p(doc, "Không được dùng tọa độ marker theo centimet rồi dùng mesh theo mét. Tất cả object_points_world, camera_position, mesh vertices và fire_xyz_world phải cùng hệ trục.")
    add_heading(doc, "Đặt marker ArUco", 2)
    add_p(doc, "Không đặt tất cả marker thành một cụm phẳng nhỏ. Hãy trải marker theo chiều ngang và chiều sâu, đồng thời đưa một số marker lên độ cao khác nếu có thể. Bốn góc marker cần được đo trong cùng hệ tọa độ phòng.")
    add_table(doc, ["ID", "Vị trí minh họa", "Cách dùng"], [
        ["0–3", "vùng thấp phía trước hoặc tường trước", "Tạo độ phủ ngang"],
        ["4–5", "hai phía trái và phải", "Mở rộng vùng nhìn"],
        ["6–7", "vùng xa hoặc độ cao khác", "Tăng điều kiện 3D cho PnP"],
    ], widths=[2.0, 7.0, 8.0])
    add_p(doc, "Các tọa độ minh họa không được dùng thay cho số đo thật. Khi dán xong, ghi ID, tọa độ bốn góc, mặt phẳng đặt marker và ngày đo. Sau khi solvePnP không di chuyển marker hoặc tripod.")

    add_heading(doc, "Thiết lập camera iPhone", 1)
    add_bullets(doc, [
        "Dùng camera chính 1x; không dùng ultra-wide hoặc digital zoom.",
        "Chọn một độ phân giải cố định, khuyến nghị 1920 × 1080 cho lần đầu.",
        "Giữ nguyên lens, độ phân giải, FPS và hướng camera trong toàn bộ session.",
        "Khóa focus và exposure; tránh autofocus hoặc auto exposure thay đổi giữa các frame.",
        "Đặt tripod chắc chắn; không tháo điện thoại ra giữa calibration và test.",
        "Tránh nắng trực tiếp, phản chiếu kính/gương và người đi qua vùng marker.",
    ])
    add_heading(doc, "Chụp checkerboard để lấy camera nội tại", 1)
    add_p(doc, "Chụp 15–25 ảnh checkerboard ở nhiều vị trí trong toàn khung hình. Bảng phải nằm phẳng, có đủ góc và không bị lóa. Cần có ảnh ở giữa, bốn góc, gần, xa và nhiều góc nghiêng.")
    add_table(doc, ["Thông số", "Giá trị"], [
        ["OpenCV pattern_size", "(9, 6)"],
        ["Số ô thực tế", "10 × 7"],
        ["square_size", "0.025 m nếu đo đúng 25 mm"],
        ["Số ảnh", "15–25"],
        ["Đầu ra", "K, dist_coeffs, image_size, reprojection error"],
    ], widths=[7.0, 10.0])
    add_p(doc, "Kết quả cần lưu theo dạng camera_intrinsics.json. Không dùng camera_calibration.example.json cho kết quả báo cáo vì đó chỉ là template.")

    doc.add_page_break()
    add_heading(doc, "Xác định pose camera bằng ArUco và solvePnP", 1)
    add_p(doc, "Sau intrinsic calibration, chụp một hoặc vài ảnh có các marker phòng. Từ mỗi marker, lấy bốn pixel corner theo thứ tự top-left, top-right, bottom-right, bottom-left. Ghép từng pixel corner với tọa độ 3D tương ứng đã đo.")
    add_code(doc, "top-left -> top-right -> bottom-right -> bottom-left")
    add_p(doc, "Nên dùng tối thiểu 6 marker nhìn thấy, tương đương 24 cặp góc 2D–3D. Bốn điểm đồng phẳng là mức tối thiểu toán học nhưng dễ bất ổn và không nên dùng làm cấu hình chính.")
    add_heading(doc, "Cấu trúc correspondences", 2)
    add_code(doc, "{\n  \"image_size\": [1920, 1080],\n  \"object_points_world\": [[x0, y0, z0], [x1, y1, z1]],\n  \"image_points\": [[u0, v0], [u1, v1]],\n  \"intrinsics\": {\n    \"K\": [[fx, 0, cx], [0, fy, cy], [0, 0, 1]],\n    \"dist_coeffs\": [k1, k2, p1, p2, k3]\n  }\n}")
    add_heading(doc, "Chạy solvePnP", 2)
    add_code(doc, ".\\.venv\\Scripts\\python.exe calibrate_room_pose_pnp.py `\n  --correspondences working\\room_marker_correspondences.json `\n  --output working\\measured_camera.json `\n  --status measured_real_room")
    add_p(doc, "Nếu intrinsic được lưu riêng, thêm --intrinsics working\\camera_intrinsics.json. Kiểm tra mean, median và P95 reprojection error. Median dưới khoảng 2 px là mục tiêu tốt cho prototype; nếu P95 vượt 5–10 px, kiểm tra lại thứ tự góc, kích thước marker, distortion và pixel corner.")
    add_heading(doc, "Các file sau bước calibration", 2)
    add_table(doc, ["File", "Nội dung", "Trạng thái"], [
        ["camera_intrinsics.json", "K, distortion, image_size", "Đo từ checkerboard"],
        ["room_marker_correspondences.json", "object_points_world và image_points", "Đo từ ArUco"],
        ["measured_camera.json", "K, R, t, camera_position, reprojection error", "Sinh từ solvePnP"],
    ], widths=[6.0, 7.0, 4.0])

    add_heading(doc, "Dựng mesh phòng theo mét", 1)
    add_p(doc, "Lần đầu chỉ cần mesh mặt sàn hình chữ nhật. Đo bốn góc vùng test theo hệ tọa độ đã chọn và chia thành hai tam giác. Sau khi sàn chạy ổn, thêm mặt bàn, nóc tủ, cột hoặc tường dưới dạng các mặt phẳng/khối có tọa độ đo thật.")
    add_code(doc, "{\n  \"metadata\": {\"status\": \"measured_real_room\", \"units\": \"metres\"},\n  \"vertices\": [[0,0,0], [4,0,0], [4,6,0], [0,6,0]],\n  \"faces\": [[0,1,2], [0,2,3]]\n}")
    add_p(doc, "Không dùng measured_room_mesh.example.json để báo cáo. File mẫu chỉ giúp hiểu schema; cần thay vertices/faces bằng số đo của phòng.")

    doc.add_page_break()
    add_heading(doc, "Tạo ground truth XYZ bằng LED", 1)
    add_p(doc, "Đặt LED tại nhiều vị trí đã đo. Khuyến nghị tối thiểu 9 vị trí trên sàn theo lưới gần–giữa–xa và trái–giữa–phải. Sau đó thêm một số vị trí trên mặt bàn nếu muốn kiểm tra multi-plane.")
    add_table(doc, ["Nhóm vị trí", "Ví dụ", "Mục đích"], [
        ["Sàn gần", "X = 0.5–2.5, Y = 0.8", "Kiểm tra ray ở khoảng cách gần"],
        ["Sàn giữa", "X = 0.5–2.5, Y = 2.0", "Kiểm tra vùng trung tâm"],
        ["Sàn xa", "X = 0.5–2.5, Y = 3.2", "Đo khuếch đại sai số pixel"],
        ["Mặt bàn", "Z = chiều cao bàn đo thật", "Kiểm tra plane selection và mesh"],
    ], widths=[4.0, 7.0, 6.0])
    add_p(doc, "Phải thống nhất điểm đại diện của LED: dùng tâm LED hoặc điểm đáy của vùng sáng. Với bài toán chân lửa, nên ghi rõ quy ước trong metadata và dùng quy ước đó cho mọi frame.")
    add_heading(doc, "Annotation tối thiểu", 2)
    add_code(doc, "{\n  \"image_path\": \"test/position_01_frame_000.jpg\",\n  \"sequence_id\": \"position_01\",\n  \"p_fire_pixel\": [812.4, 646.1],\n  \"fire_xyz_world\": [1.25, 3.40, 0.00],\n  \"source_type\": \"led\",\n  \"marker_id\": \"led_01\"\n}")
    add_p(doc, "Tách các vị trí hoặc sequence giữa train, validation và test. Không để các frame liên tiếp cùng một vị trí xuất hiện ở cả train và test.")
    add_heading(doc, "Tạo manifest dữ liệu thật", 2)
    add_code(doc, ".\\.venv\\Scripts\\python.exe build_real_flame_manifest.py `\n  --dataset-root working\\iphone_room_session `\n  --annotations working\\iphone_room_annotations.json `\n  --output working\\iphone_room_session\\manifest.jsonl `\n  --require-xyz --source-type led")

    add_heading(doc, "Chạy Ray Casting", 1)
    add_p(doc, "Dùng ROI mixed hiện tại cho điểm 2D trước khi chiếu tia. Checkpoint khuyến nghị là output\\roi_domain_experiments_cpu_regularized\\mixed\\best_roi.pth. Ray Casting vật lý phải dùng measured_camera.json và measured_room_mesh.json, không dùng file example.")
    add_code(doc, ".\\.venv\\Scripts\\python.exe paper_workflow_3d.py `\n  --dataset working\\iphone_room_session `\n  --coarse-source detector `\n  --detector-checkpoint fire-model-data\\best.pth `\n  --roi-checkpoint output\\roi_domain_experiments_cpu_regularized\\mixed\\best_roi.pth `\n  --calibration-source external `\n  --calibration working\\measured_camera.json `\n  --mesh working\\measured_room_mesh.json `\n  --labels-3d working\\measured_ground_truth_3d.json `\n  --output-dir output\\physical_raycast_first `\n  --split test --max-records 0 --device cpu")
    add_p(doc, "Nếu chỉ muốn kiểm tra hình học bằng pixel đã có trong manifest, có thể dùng coarse-source manifest. Khi đánh giá detector end-to-end, dùng detector checkpoint và ghi riêng các trường hợp detector miss.")

    add_heading(doc, "Cấu hình multi ray và temporal filter", 2)
    add_table(doc, ["Thử nghiệm", "Thiết lập", "Mục tiêu"], [
        ["Single ray", "--ray-pattern single", "Baseline tối giản"],
        ["3 rays", "cross với bán kính nhỏ", "Giảm nhạy với một pixel lỗi"],
        ["5 rays", "cross/square tùy adapter", "Kiểm tra độ ổn định trung tâm"],
        ["7 rays", "pattern square hoặc cấu hình mở rộng", "Đánh giá chi phí và hit rate"],
        ["EMA", "--temporal-filter ema", "Giảm jitter trên video"],
        ["EKF", "--temporal-filter ekf", "Theo dõi vị trí có vận tốc"],
    ], widths=[4.0, 7.0, 6.0])
    add_p(doc, "Không dùng EMA/EKF để làm đẹp kết quả sai. Trước khi lọc, phải ghi ray hit/miss, khoảng cách nhảy và lý do fallback. Smoothing chỉ áp dụng sau khi đã loại ray không hợp lệ và outlier rõ ràng.")

    doc.add_page_break()
    add_heading(doc, "Benchmark và tiêu chí nghiệm thu", 1)
    add_p(doc, "Benchmark cuối phải chạy trên cùng camera, cùng mesh, cùng tập test và cùng ground truth XYZ. Không gộp số liệu synthetic, asset-backed và measured real-room vào một bảng.")
    add_table(doc, ["Nhóm chỉ số", "Cần báo cáo"], [
        ["2D", "Pixel MAE, median, P95, PCK@10, PCK@25"],
        ["Ray geometry", "Ray hit rate, valid multi-ray rate, fallback rate"],
        ["3D tổng", "MAE, median, P95 theo mét"],
        ["Theo trục", "MAE và P95 riêng cho X, Y, Z"],
        ["Ngưỡng", "Tỷ lệ sai số dưới 0.10 m, 0.25 m, 0.50 m, 1.00 m"],
        ["Video", "Jitter, jump rate, số frame mất track, EMA/EKF stability"],
        ["Hiệu năng", "Latency mỗi ảnh/frame và số frame xử lý"],
        ["Calibration", "Reprojection mean, median, P95 theo pixel"],
    ], widths=[5.0, 12.0])
    add_heading(doc, "Xuất ảnh 3D", 1)
    add_code(doc, ".\\.venv\\Scripts\\python.exe visualize_3d_results.py `\n  --summary output\\physical_raycast_first\\summary.json `\n  --dataset working\\iphone_room_session `\n  --mesh working\\measured_room_mesh.json `\n  --calibration working\\measured_camera.json `\n  --calibration-source external `\n  --output-dir output\\physical_raycast_first\\visualization `\n  --write-ply")
    add_p(doc, "Các file cần kiểm tra sau khi xuất gồm scene_3d_overview.png, scene_3d_top_view.png, scene_3d_side_view.png, scene_3d_interactive.html, scene_3d_annotations.ply và visualization_manifest.json.")
    add_heading(doc, "Kiểm tra trực quan", 2)
    add_checklist(doc, [
        "Mesh phòng nằm đúng kích thước và hướng trục",
        "Camera xuất hiện ở vị trí hợp lý",
        "Ground truth LED và điểm Ray Casting nằm trong cùng vùng phòng",
        "Ray không xuyên sai qua tường hoặc mặt bàn",
        "Các màu GT, coarse, ROI, Ray và filtered được phân biệt rõ",
        "HTML có thể xoay, phóng to và kiểm tra từng điểm",
        "PLY mở được trong CloudCompare hoặc MeshLab",
    ])

    add_heading(doc, "Lỗi thường gặp", 1)
    add_table(doc, ["Hiện tượng", "Nguyên nhân thường gặp", "Cách xử lý"], [
        ["PnP error lớn", "Sai thứ tự góc, marker cong, điểm 3D sai", "Kiểm tra TL-TR-BR-BL, đo lại và dùng 6–10 marker"],
        ["Ray miss nhiều", "Mesh thiếu vùng sàn, K/R sai hoặc pixel gần horizon", "Kiểm tra cùng hệ mét và reprojection trước"],
        ["Điểm 3D lệch đều", "Sai pose hoặc sai gốc tọa độ", "Kiểm tra camera_position, trục và mesh"],
        ["Điểm nhảy mạnh theo video", "ROI/coarse outlier hoặc calibration nhiễu", "Ghi raw trước, rồi dùng outlier gate và EMA/EKF"],
        ["Marker không detect", "In nhỏ, lóa, dictionary không khớp", "Dùng đúng DICT_6X6_1000, tăng kích thước và ánh sáng đều"],
        ["Kết quả đẹp nhưng sai mét", "Dùng file example hoặc mesh chưa đo", "Chỉ báo cáo physical accuracy với dữ liệu measured"],
    ], widths=[4.0, 7.0, 6.0])

    doc.add_page_break()
    add_heading(doc, "Checklist thực hiện tại phòng", 1)
    add_p(doc, "In trang này hoặc dùng như danh sách nghiệm thu trước khi chạy benchmark.")
    add_checklist(doc, [
        "Đã chọn vùng test 3–4 m × 4–6 m và dọn các vật thể chuyển động",
        "Đã đánh dấu gốc O, trục X, trục Y trên sàn",
        "Đã đo kích thước vùng test bằng mét",
        "Đã cố định iPhone trên tripod và khóa camera settings",
        "Đã chụp 15–25 ảnh checkerboard",
        "Đã tạo camera_intrinsics.json từ checkerboard",
        "Đã đặt và đo tọa độ 6–10 marker ArUco",
        "Đã chụp ảnh marker và tạo correspondences 2D–3D",
        "Đã chạy solvePnP và kiểm tra reprojection error",
        "Đã tạo measured_camera.json với status measured_real_room",
        "Đã tạo measured_room_mesh.json cùng hệ tọa độ",
        "Đã đặt LED tại tối thiểu 9 vị trí và đo fire_xyz_world",
        "Đã chụp test ảnh/video riêng với train/validation",
        "Đã tạo manifest có p_fire_pixel và fire_xyz_world",
        "Đã chạy baseline single-ray trước multi-ray",
        "Đã chạy 3/5/7 rays và ghi hit rate, latency",
        "Đã chạy EMA/EKF trên video theo thứ tự frame",
        "Đã xuất PNG, HTML và PLY",
        "Đã lưu raw predictions, fallback, miss và uncertainty",
        "Đã tách kết quả physical khỏi synthetic/provisional",
    ])
    add_heading(doc, "Kết luận sử dụng", 1)
    add_p(doc, "Một thí nghiệm Ray Casting vật lý có giá trị khi camera calibration, pose camera, mesh và ground truth XYZ đều được đo trong cùng một hệ tọa độ. Với thiết bị hiện có là iPhone, tripod và LED, phần còn lại chủ yếu là in target, đo đạc cẩn thận và ghi nhãn nhất quán. Không cần mua LiDAR, camera stereo hoặc tạo lửa thật ở giai đoạn đầu.")
    add_p(doc, "Checkpoint ROI mixed chỉ cung cấp điểm 2D. Nó không thay thế calibration hoặc mesh. Vì vậy hãy nghiệm thu theo thứ tự: calibration → sàn đơn giản → single-ray → multi-ray → temporal filter → thêm vật cản → benchmark cuối.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
